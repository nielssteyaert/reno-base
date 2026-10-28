"""Spreadsheet import: read an uploaded .xlsx/.csv, guess the column mapping, preview, commit.

Uploads live under <data>/imports/<token><ext> (plus <token>.name holding the original file
name) between the steps of the import wizard. They are removed after a successful commit, and
any upload older than 24 hours is removed on the next upload.

Expenses are created through app.services.expenses.create_expense, so VAT is split exactly as
in the UI and the audit hook records every inserted row.
"""

from __future__ import annotations

import csv
import io
import re
import time
import unicodedata
import uuid
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from fastapi import HTTPException, UploadFile
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app import schemas
from app.config import get_config
from app.models import Category, Payment, PaymentMethod, Room, Stage, Supplier
from app.money import amounts
from app.services import excel_export
from app.services.budget import get_settings
from app.services.expenses import create_expense

MAX_BYTES = 10 * 1024 * 1024
MAX_AGE_SECONDS = 24 * 3600
EXTENSIONS = (".xlsx", ".csv")
SAMPLE_ROWS = 10
_CHUNK = 1024 * 1024
_TOKEN_RE = re.compile(r"[0-9a-f]{32}")
_PREFERRED_SHEETS = ("expenses", "uitgaven", "kosten")


# --- Target fields ------------------------------------------------------------------------


@dataclass(frozen=True)
class ImportField:
    key: str
    label: str
    kind: str  # text, amount, percent, date, stage, ref
    required: bool = False
    # Header names recognised by guess_mapping, best first (English and Dutch).
    synonyms: tuple[str, ...] = ()
    help: str = ""


# fmt: off
FIELDS: tuple[ImportField, ...] = (
    ImportField(
        "title", "Title", "text", True,
        ("title", "titel", "naam", "name", "item", "omschrijving", "post", "artikel",
         "product", "onderwerp", "uitgave", "expense", "beschrijving", "description"),
        "Short name of the expense.",
    ),
    ImportField(
        "amount", "Amount", "amount", True,
        ("amount", "bedrag", "prijs", "price", "totaal", "total", "kost", "kosten", "cost",
         "budget", "excl. VAT", "amount excl. VAT", "bedrag excl. btw", "prijs excl. btw",
         "excl. btw", "incl. VAT", "amount incl. VAT", "bedrag incl. btw", "prijs incl. btw",
         "incl. btw", "totaal incl. btw"),
        "Excl. VAT, or incl. VAT when 'amounts include VAT' is chosen while importing. "
        "Numbers, or text like 1.234,56 / 1,234.56 / € 1 234,50.",
    ),
    ImportField(
        "vat_rate", "VAT %", "percent", False,
        ("VAT %", "BTW %", "VAT rate", "BTW tarief", "BTW percentage", "tarief", "VAT", "BTW",
         "tax rate", "tax"),
        "VAT percentage: 21, 21%, 0.21 or 6,0. Empty uses the default VAT rate.",
    ),
    ImportField(
        "category", "Category", "ref", False,
        ("category", "categorie", "hoofdcategorie", "rubriek", "groep", "group", "soort", "type"),
        "Category name. 'Parent > Child' or 'Parent / Child' also sets the subcategory. "
        "Unknown names can be linked to an existing category or created.",
    ),
    ImportField(
        "subcategory", "Subcategory", "ref", False,
        ("subcategory", "subcategorie", "sub category", "sub categorie", "subcat", "subgroep"),
        "Subcategory under the category; created automatically when it does not exist.",
    ),
    ImportField(
        "room", "Room", "ref", False,
        ("room", "ruimte", "kamer", "locatie", "location", "plaats", "zone", "area"),
        "Room name. Unknown names can be linked to an existing room or created.",
    ),
    ImportField(
        "supplier", "Supplier", "ref", False,
        ("supplier", "leverancier", "aannemer", "contractor", "vendor", "firma", "bedrijf",
         "company", "vakman", "winkel", "shop"),
        "Supplier or contractor name. Unknown names can be linked to an existing one or created.",
    ),
    ImportField(
        "stage", "Status", "stage", False,
        ("status", "stage", "fase", "phase", "stadium"),
        "estimate (raming, schatting, budget), quote (offerte), ordered (besteld, akkoord), "
        "invoiced (factuur, gefactureerd) or paid (betaald: invoiced and fully paid). "
        "Empty uses the default status.",
    ),
    ImportField(
        "expense_date", "Date", "date", False,
        ("date", "datum", "expense date", "factuurdatum", "invoice date", "datum factuur",
         "besteldatum", "order date"),
        "dd/mm/yyyy, dd-mm-yyyy, dd.mm.yyyy or yyyy-mm-dd (2-digit years allowed), or a date cell.",
    ),
    ImportField(
        "due_date", "Due date", "date", False,
        ("due date", "due", "vervaldatum", "vervaldag", "deadline", "payment due",
         "uiterste betaaldatum"),
        "When payment is due. Same formats as Date.",
    ),
    ImportField(
        "quote_ref", "Quote ref", "text", False,
        ("quote ref", "quote", "quote number", "quote no", "offerte", "offertenummer",
         "offerte nr", "offerte ref", "referentie offerte"),
        "Quote reference (free text).",
    ),
    ImportField(
        "invoice_no", "Invoice no", "text", False,
        ("invoice no", "invoice", "invoice number", "factuur", "factuurnummer", "factuur nr",
         "factuurnr"),
        "Invoice number (free text).",
    ),
    ImportField(
        "description", "Description", "text", False,
        ("description", "beschrijving", "omschrijving", "details", "detail", "specificatie"),
        "Longer description (free text).",
    ),
    ImportField(
        "notes", "Notes", "text", False,
        ("notes", "note", "notities", "notitie", "opmerking", "opmerkingen", "comment",
         "comments", "remarks", "memo"),
        "Notes (free text).",
    ),
    ImportField(
        "paid_amount", "Paid amount", "amount", False,
        ("paid amount", "paid", "betaald", "betaald bedrag", "amount paid", "voldaan",
         "reeds betaald", "al betaald", "payment", "betaling"),
        "Amount already paid, incl. VAT. Creates one bank-transfer payment.",
    ),
    ImportField(
        "paid_date", "Paid date", "date", False,
        ("paid date", "betaaldatum", "payment date", "datum betaling", "betaald op", "paid on",
         "betalingsdatum"),
        "Date of that payment. Empty uses the expense date, or today.",
    ),
)
# fmt: on
FIELD_KEYS = tuple(f.key for f in FIELDS)
REF_KINDS = ("category", "room", "supplier")


def field_list() -> list[dict]:
    return [
        {"key": f.key, "label": f.label, "required": f.required, "kind": f.kind} for f in FIELDS
    ]


# --- Header matching ----------------------------------------------------------------------


def normalize(text: Any) -> str:
    """Lowercase, accent-free, punctuation-free form of a header or name."""
    decomposed = unicodedata.normalize("NFKD", str(text))
    plain = "".join(c for c in decomposed if not unicodedata.combining(c)).lower()
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", plain.replace("%", " % ")).split())


_SYNONYMS = {
    f.key: tuple(dict.fromkeys(normalize(s) for s in (f.label, *f.synonyms))) for f in FIELDS
}
# Partial matches ("Bedrag excl. btw (€)" contains "bedrag") only use specific words.
_MIN_PARTIAL = 5


def guess_mapping(headers: list[str]) -> dict[str, int | None]:
    """Map each field to the column whose header looks like it; each column is used once."""
    names = [normalize(h) for h in headers]
    mapping: dict[str, int | None] = dict.fromkeys(FIELD_KEYS)
    used: set[int] = set()
    # Exact names first, in field order, so "omschrijving" becomes the title when there is no
    # better title column and the description otherwise.
    for key in FIELD_KEYS:
        for synonym in _SYNONYMS[key]:
            col = next((i for i, n in enumerate(names) if n == synonym and i not in used), None)
            if col is not None:
                mapping[key] = col
                used.add(col)
                break
    candidates = []
    for order, key in enumerate(FIELD_KEYS):
        if mapping[key] is not None:
            continue
        for synonym in _SYNONYMS[key]:
            if len(synonym) < _MIN_PARTIAL:
                continue
            for col, name in enumerate(names):
                if col not in used and f" {synonym} " in f" {name} ":
                    candidates.append((-len(synonym), order, col, key))
    for _, _, col, key in sorted(candidates):
        if mapping[key] is None and col not in used:
            mapping[key] = col
            used.add(col)
    return mapping


# --- Value parsing ------------------------------------------------------------------------


def _round_cents(value: Decimal) -> int:
    return int(value.quantize(Decimal(1), rounding=ROUND_HALF_UP))


def parse_number(text: str) -> Decimal | None:
    """Parse a number written with either decimal separator; None when empty.

    The last "," or "." is the decimal separator when both kinds occur, or when it occurs
    once and is not followed by exactly three digits (then it groups thousands).
    """
    text = re.sub(r"(?i)eur(o)?|[€$£\s  '’]", "", text)
    if text in ("", "-", "–", "—"):
        return None
    negative = text.startswith(("-", "−")) or (text.startswith("(") and text.endswith(")"))
    text = text.strip("()+-−")
    if not re.fullmatch(r"[\d.,]*\d[\d.,]*", text):
        raise ValueError(text)
    last = max(text.rfind(","), text.rfind("."))
    if last == -1:
        number = Decimal(text)
    else:
        sep = text[last]
        tail = text[last + 1 :]
        other = "." if sep == "," else ","
        is_decimal = other in text or (text.count(sep) == 1 and len(tail) != 3)
        if is_decimal:
            whole = re.sub(r"[.,]", "", text[:last]) or "0"
            number = Decimal(f"{whole}.{tail or 0}")
        else:
            number = Decimal(re.sub(r"[.,]", "", text))
    return -number if negative else number


def parse_amount(value: Any) -> int | None:
    """Amount in cents from a cell; None when empty. Raises ValueError when invalid/negative."""
    if value is None:
        return None
    if isinstance(value, bool):
        raise ValueError(value)
    if isinstance(value, int | float | Decimal):
        number = Decimal(str(value))
    else:
        number = parse_number(str(value))
        if number is None:
            return None
    try:
        cents = _round_cents(number * 100)
    except InvalidOperation as exc:
        raise ValueError(value) from exc
    if cents < 0:
        raise ValueError("negative")
    return cents


def parse_vat(value: Any) -> float | None:
    """VAT percentage from 21, "21%", 0.21 or "6,0"; None when empty."""
    if value is None:
        return None
    if isinstance(value, bool):
        raise ValueError(value)
    percent_sign = False
    if isinstance(value, int | float | Decimal):
        rate = float(value)
    else:
        text = str(value).strip()
        percent_sign = "%" in text
        number = parse_number(text.replace("%", ""))
        if number is None:
            return None
        rate = float(number)
    if not percent_sign and 0 < rate <= 1:
        rate *= 100
    rate = round(rate, 2)
    if not 0 <= rate <= 100:
        raise ValueError(value)
    return rate


_DMY = re.compile(r"(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{4}|\d{2})")
_YMD = re.compile(r"(\d{4})[/.\-](\d{1,2})[/.\-](\d{1,2})")


def parse_date(value: Any) -> date | None:
    """Date from a date cell, an Excel serial number or day-first text; None when empty."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, int | float) and not isinstance(value, bool):
        if 1 <= value < 2_958_466:  # Excel's serial range (up to 31/12/9999)
            return from_excel(value).date()
        raise ValueError(value)
    text = str(value).strip().split("T")[0].split(" ")[0]
    if match := _YMD.fullmatch(text):
        year, month, day = map(int, match.groups())
    elif match := _DMY.fullmatch(text):
        day, month, year = map(int, match.groups())
        if year < 100:
            year += 2000 if year < 70 else 1900
    else:
        raise ValueError(value)
    return date(year, month, day)


_STAGE_WORDS = {
    **dict.fromkeys(
        ("estimate", "estimated", "raming", "schatting", "geschat", "budget"), Stage.ESTIMATE
    ),
    **dict.fromkeys(("quote", "quoted", "quotation", "offerte"), Stage.QUOTE),
    **dict.fromkeys(("ordered", "order", "besteld", "bestelling", "akkoord"), Stage.ORDERED),
    **dict.fromkeys(("invoiced", "invoice", "factuur", "gefactureerd"), Stage.INVOICED),
}
_PAID_WORDS = {"paid", "betaald", "voldaan"}


def parse_stage(value: Any) -> tuple[Stage | None, bool]:
    """(stage, fully_paid) for a status word; (None, False) when empty.

    "paid"/"betaald" means invoiced and fully paid. Raises ValueError for an unknown word.
    """
    words = normalize(value).split() if value is not None else []
    if not words:
        return None, False
    for candidate in (" ".join(words), words[0]):
        if candidate in _PAID_WORDS:
            return Stage.INVOICED, True
        if candidate in _STAGE_WORDS:
            return _STAGE_WORDS[candidate], False
    raise ValueError(value)


def split_category(value: str | None) -> tuple[str | None, str | None]:
    """'Parent > Child' or 'Parent / Child' into (parent, child)."""
    if not value:
        return None, None
    parts = re.split(r"\s*[>/]\s*", value, maxsplit=1)
    parent = parts[0].strip() or None
    child = parts[1].strip() if len(parts) > 1 else ""
    return parent, child or None


def _text(value: Any) -> str | None:
    """Cell value as stripped text; None when empty."""
    if value is None:
        return None
    if isinstance(value, datetime):
        value = value.date() if value.time() == datetime.min.time() else value
    if isinstance(value, date) and not isinstance(value, datetime):
        return f"{value:%d/%m/%Y}"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


# --- Reading uploads ----------------------------------------------------------------------


@dataclass
class Table:
    sheets: list[str]
    sheet: str | None
    headers: list[str]
    rows: list[list[Any]]


def imports_dir() -> Path:
    return get_config().data_dir / "imports"


def _clean(value: Any) -> Any:
    if isinstance(value, str):
        return value.strip() or None
    return value


def _split_header(raw_rows: list[list[Any]]) -> tuple[list[str], list[list[Any]]]:
    """The first non-empty row is the header; fully empty rows are dropped."""
    rows = [[_clean(v) for v in row] for row in raw_rows]
    rows = [row for row in rows if any(v is not None for v in row)]
    if not rows:
        return [], []
    header = rows[0]
    width = max(i for i, v in enumerate(header) if v is not None) + 1
    headers = [
        _text(v) or f"Column {get_column_letter(i + 1)}" for i, v in enumerate(header[:width])
    ]
    data = []
    for row in rows[1:]:
        row = (row + [None] * width)[:width]
        if any(v is not None for v in row):
            data.append(row)
    return headers, data


def _sniff_delimiter(text: str) -> str:
    """The delimiter that splits the header into most columns consistently across rows."""
    sample = "\n".join([line for line in text.splitlines() if line.strip()][:20])
    best, best_score = ";", (0, 0)
    for delimiter in (";", ",", "\t"):
        rows = list(csv.reader(io.StringIO(sample), delimiter=delimiter))
        width = len(rows[0]) if rows else 0
        if width < 2:
            continue
        score = (sum(1 for r in rows if len(r) == width), width)
        if score > best_score:
            best, best_score = delimiter, score
    return best


def _read_csv(path: Path) -> Table:
    data = path.read_bytes()
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252", errors="replace")
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=_sniff_delimiter(text))
    headers, rows = _split_header([list(row) for row in reader])
    return Table([], None, headers, rows)


def _read_xlsx(path: Path, sheet: str | None) -> Table:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        names = [ws.title for ws in wb.worksheets]
        if not names:
            raise HTTPException(422, "The workbook has no sheets")
        if sheet is None:
            sheet = next((n for n in names if n.lower() in _PREFERRED_SHEETS), names[0])
        elif sheet not in names:
            raise HTTPException(422, f"Sheet '{sheet}' not found")
        ws = wb[sheet]
        # Some writers store a wrong sheet size; recompute it so no rows are missed.
        ws.reset_dimensions()
        raw = [list(row) for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    headers, rows = _split_header(raw)
    return Table(names, sheet, headers, rows)


def read_table(path: Path, sheet: str | None = None) -> Table:
    if path.suffix == ".csv":
        return _read_csv(path)
    return _read_xlsx(path, sheet)


def upload_path(token: str) -> Path:
    """The stored upload for a token; 404 when unknown or expired."""
    if _TOKEN_RE.fullmatch(token):
        for suffix in EXTENSIONS:
            path = imports_dir() / f"{token}{suffix}"
            if path.is_file():
                return path
    raise HTTPException(404, "Upload not found or expired; please upload the file again")


def remove_upload(path: Path) -> None:
    path.unlink(missing_ok=True)
    path.with_suffix(".name").unlink(missing_ok=True)


def _remove_stale(directory: Path) -> None:
    cutoff = time.time() - MAX_AGE_SECONDS
    for path in directory.iterdir():
        if path.is_file() and path.stat().st_mtime < cutoff:
            path.unlink(missing_ok=True)


async def save_upload(upload: UploadFile) -> str:
    """Store an uploaded .xlsx/.csv (max 10 MB) and return its token."""
    original = upload.filename or "upload"
    suffix = Path(original).suffix.lower()
    if suffix not in EXTENSIONS:
        raise HTTPException(415, "Only .xlsx and .csv files can be imported")
    directory = imports_dir()
    directory.mkdir(parents=True, exist_ok=True)
    _remove_stale(directory)
    token = uuid.uuid4().hex
    target = directory / f"{token}{suffix}"
    size = 0
    try:
        with target.open("wb") as fh:
            while chunk := await upload.read(_CHUNK):
                size += len(chunk)
                if size > MAX_BYTES:
                    raise HTTPException(413, "File is larger than 10 MB")
                fh.write(chunk)
    except BaseException:
        target.unlink(missing_ok=True)
        raise
    target.with_suffix(".name").write_text(original[:255], encoding="utf-8")
    return token


def _display(value: Any) -> str:
    return _text(value) or ""


def inspect(token: str, sheet: str | None = None) -> dict:
    """Sheets, headers, a sample and the guessed mapping of an upload."""
    path = upload_path(token)
    try:
        table = read_table(path, sheet)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, "This file could not be read as a spreadsheet") from exc
    name_file = path.with_suffix(".name")
    return {
        "token": token,
        "filename": name_file.read_text(encoding="utf-8") if name_file.exists() else path.name,
        "sheets": table.sheets,
        "sheet": table.sheet,
        "headers": table.headers,
        "sample": [[_display(v) for v in row] for row in table.rows[:SAMPLE_ROWS]],
        "row_count": len(table.rows),
        "mapping": guess_mapping(table.headers),
    }


# --- Preview and commit -------------------------------------------------------------------


class ValueMap(BaseModel):
    category: dict[str, int | None] = Field(default_factory=dict)
    room: dict[str, int | None] = Field(default_factory=dict)
    supplier: dict[str, int | None] = Field(default_factory=dict)


class ImportRequest(BaseModel):
    token: str
    sheet: str | None = None
    mapping: dict[str, int | None]
    amount_includes_vat: bool = False
    default_vat_rate: float | None = Field(None, ge=0, le=100)
    default_stage: Stage = Stage.INVOICED
    value_map: ValueMap = Field(default_factory=ValueMap)
    skip_rows: list[int] = Field(default_factory=list)


@dataclass
class ParsedRow:
    index: int
    values: dict[str, Any]
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    entered_cents: int = 0  # the amount as typed (excl. or incl. VAT), for create_expense

    def out(self) -> dict:
        values = {k: v.isoformat() if isinstance(v, date) else v for k, v in self.values.items()}
        return {
            "index": self.index,
            "values": values,
            "errors": self.errors,
            "warnings": self.warnings,
        }


def _key(name: str) -> str:
    return name.strip().casefold()


_MODELS = {"category": Category, "room": Room, "supplier": Supplier}
_NAME_LIMITS = {"category": 120, "room": 120, "supplier": 200}


class _Names:
    """Existing categories, rooms and suppliers by case-insensitive name, plus the value map.

    Names created during a commit are added here, so each new name is created only once.
    """

    def __init__(self, db: Session, value_map: ValueMap):
        self.db = db
        self.existing: dict[str, dict[str, Any]] = {kind: {} for kind in REF_KINDS}
        self.children: dict[tuple[int, str], Category] = {}
        for category in db.scalars(select(Category).order_by(Category.id)).all():
            if category.parent_id is None:
                self.existing["category"].setdefault(_key(category.name), category)
            else:
                self.children.setdefault((category.parent_id, _key(category.name)), category)
        for kind in ("room", "supplier"):
            for entity in db.scalars(select(_MODELS[kind]).order_by(_MODELS[kind].id)).all():
                self.existing[kind].setdefault(_key(entity.name), entity)
        self.value_map: dict[str, dict[str, int | None]] = {}
        for kind in REF_KINDS:
            entries = getattr(value_map, kind)
            for entity_id in entries.values():
                if entity_id is not None and db.get(_MODELS[kind], entity_id) is None:
                    raise HTTPException(422, f"Unknown {kind} id {entity_id} in value_map")
            self.value_map[kind] = {_key(name): i for name, i in entries.items()}
        self.created = dict.fromkeys(REF_KINDS, 0)

    def is_unmatched(self, kind: str, name: str) -> bool:
        key = _key(name)
        return key not in self.existing[kind] and key not in self.value_map[kind]

    def mapped_subcategory(self, name: str) -> bool:
        """True when the value map links this category name to an existing subcategory."""
        entity_id = self.value_map["category"].get(_key(name))
        return entity_id is not None and self.db.get(Category, entity_id).parent_id is not None

    def resolve(self, kind: str, name: str) -> Any:
        """The record for a name: value-map link, existing record, or a newly created one."""
        key = _key(name)
        entity_id = self.value_map[kind].get(key)
        if entity_id is not None:
            return self.db.get(_MODELS[kind], entity_id)
        if key not in self.existing[kind]:
            entity = _MODELS[kind](name=name[: _NAME_LIMITS[kind]])
            self._add(kind, entity)
            self.existing[kind][key] = entity
        return self.existing[kind][key]

    def subcategory(self, parent: Category, name: str) -> Category:
        key = (parent.id, _key(name))
        if key not in self.children:
            child = Category(name=name[: _NAME_LIMITS["category"]], parent_id=parent.id)
            self._add("category", child)
            self.children[key] = child
        return self.children[key]

    def _add(self, kind: str, entity: Any) -> None:
        model = type(entity)
        if hasattr(model, "sort_order"):
            # New records go after the existing ones instead of jumping to the top.
            entity.sort_order = (self.db.scalar(select(func.max(model.sort_order))) or 0) + 1
        self.db.add(entity)
        self.db.flush()  # assigns the id; the whole import is still one transaction
        self.created[kind] += 1


def _check_mapping(request: ImportRequest, width: int) -> dict[str, int]:
    mapping = {k: v for k, v in request.mapping.items() if k in FIELD_KEYS and v is not None}
    missing = [f.label for f in FIELDS if f.required and f.key not in mapping]
    if missing:
        raise HTTPException(422, f"Map a column to: {', '.join(missing)}")
    for key, col in mapping.items():
        if not 0 <= col < width:
            raise HTTPException(422, f"Column {col} for '{key}' does not exist")
    return mapping


def _parse_row(
    index: int, raw: list[Any], mapping: dict[str, int], request: ImportRequest, vat_default: float
) -> ParsedRow:
    def cell(key: str) -> Any:
        return raw[mapping[key]] if key in mapping else None

    def text(key: str, limit: int | None = None) -> str | None:
        value = _text(cell(key))
        return value[:limit] if value and limit else value

    row = ParsedRow(index, {})
    errors, warnings = row.errors, row.warnings

    title = text("title")
    if not title:
        errors.append("Title is missing")
    elif len(title) > 200:
        warnings.append("Title shortened to 200 characters")
        title = title[:200]

    amount = None
    try:
        amount = parse_amount(cell("amount"))
        if amount is None:
            errors.append("Amount is missing")
    except ValueError:
        errors.append(f"Amount '{_display(cell('amount'))}' is not a valid amount (0 or more)")

    vat_rate = vat_default
    try:
        vat_rate = parse_vat(cell("vat_rate"))
        if vat_rate is None:
            vat_rate = vat_default
    except ValueError:
        vat_rate = None
        errors.append(f"VAT '{_display(cell('vat_rate'))}' is not a valid VAT rate (0-100)")

    stage, fully_paid = request.default_stage, False
    try:
        parsed_stage, fully_paid = parse_stage(cell("stage"))
        stage = parsed_stage or request.default_stage
    except ValueError:
        warnings.append(
            f"Unknown status '{_display(cell('stage'))}', using '{request.default_stage.value}'"
        )

    dates: dict[str, date | None] = {}
    date_fields = (("expense_date", "Date"), ("due_date", "Due date"), ("paid_date", "Paid date"))
    for key, label in date_fields:
        try:
            dates[key] = parse_date(cell(key))
        except ValueError:
            dates[key] = None
            errors.append(f"{label} '{_display(cell(key))}' is not a valid date (dd/mm/yyyy)")

    paid_amount = None
    try:
        paid_amount = parse_amount(cell("paid_amount"))
    except ValueError:
        errors.append(f"Paid amount '{_display(cell('paid_amount'))}' is not a valid amount")

    excl = vat = incl = None
    if amount is not None and vat_rate is not None:
        excl, vat = amounts(amount, vat_rate, request.amount_includes_vat)
        incl = excl + vat
        row.entered_cents = amount
    if paid_amount is not None:
        paid = paid_amount
    else:
        paid = (incl or 0) if fully_paid else 0
    if incl is not None and paid > incl:
        warnings.append("Paid amount is more than the amount incl. VAT")
    paid_date = dates["paid_date"]
    if paid and paid_date is None:
        paid_date = dates["expense_date"] or date.today()

    category, subcategory = split_category(text("category"))
    subcategory = text("subcategory") or subcategory
    if subcategory and not category:
        warnings.append(f"Subcategory '{subcategory}' ignored: no category given")
        subcategory = None

    row.values = {
        "title": title,
        "amount_excl_cents": excl,
        "vat_rate": vat_rate,
        "vat_cents": vat,
        "amount_incl_cents": incl,
        "category": category,
        "subcategory": subcategory,
        "room": text("room"),
        "supplier": text("supplier"),
        "stage": stage.value,
        "expense_date": dates["expense_date"],
        "due_date": dates["due_date"],
        "quote_ref": text("quote_ref", 100),
        "invoice_no": text("invoice_no", 100),
        "description": text("description"),
        "notes": text("notes"),
        "paid_cents": paid,
        "paid_date": paid_date,
    }
    return row


def _prepare(db: Session, request: ImportRequest) -> tuple[list[ParsedRow], _Names, Path]:
    path = upload_path(request.token)
    table = read_table(path, request.sheet)
    mapping = _check_mapping(request, len(table.headers))
    names = _Names(db, request.value_map)
    vat_default = request.default_vat_rate
    if vat_default is None:
        vat_default = get_settings(db).default_vat_rate
    rows = [_parse_row(i, raw, mapping, request, vat_default) for i, raw in enumerate(table.rows)]
    for row in rows:
        category = row.values["category"]
        if category and row.values["subcategory"] and names.mapped_subcategory(category):
            row.warnings.append(f"Subcategory ignored: '{category}' is already a subcategory")
            row.values["subcategory"] = None
    return rows, names, path


def _importable(rows: list[ParsedRow], skip: list[int]) -> list[ParsedRow]:
    skipped = set(skip)
    return [row for row in rows if not row.errors and row.index not in skipped]


def preview(db: Session, request: ImportRequest) -> dict:
    """Parsed rows with errors/warnings, without writing anything."""
    rows, names, _ = _prepare(db, request)
    importable = _importable(rows, request.skip_rows)
    unmatched: dict[str, list[str]] = {kind: [] for kind in REF_KINDS}
    seen: set[tuple[str, str]] = set()
    for row in importable:
        for kind in REF_KINDS:
            name = row.values[kind]
            if name and (kind, _key(name)) not in seen and names.is_unmatched(kind, name):
                seen.add((kind, _key(name)))
                unmatched[kind].append(name)
    return {
        "rows": [row.out() for row in rows],
        "valid_count": len(importable),
        "error_count": sum(1 for row in rows if row.errors),
        "unmatched": unmatched,
        "total_incl_cents": sum(row.values["amount_incl_cents"] for row in importable),
    }


def commit(db: Session, request: ImportRequest) -> dict:
    """Create every importable row in one transaction, then delete the upload."""
    rows, names, path = _prepare(db, request)
    importable = _importable(rows, request.skip_rows)
    user = db.info.get("user")
    payments = 0
    try:
        for row in importable:
            values = row.values
            category = None
            if values["category"]:
                category = names.resolve("category", values["category"])
                if values["subcategory"]:
                    category = names.subcategory(category, values["subcategory"])
            room = names.resolve("room", values["room"]) if values["room"] else None
            supplier = names.resolve("supplier", values["supplier"]) if values["supplier"] else None
            payload = schemas.ExpenseCreate(
                title=values["title"],
                amount_cents=row.entered_cents,
                amount_includes_vat=request.amount_includes_vat,
                vat_rate=values["vat_rate"],
                stage=Stage(values["stage"]),
                category_id=category.id if category else None,
                room_id=room.id if room else None,
                supplier_id=supplier.id if supplier else None,
                expense_date=values["expense_date"],
                due_date=values["due_date"],
                quote_ref=values["quote_ref"],
                invoice_no=values["invoice_no"],
                description=values["description"],
                notes=values["notes"],
            )
            expense = create_expense(db, payload, user)
            if values["paid_cents"]:
                expense.payments.append(
                    Payment(
                        date=values["paid_date"],
                        amount_cents=values["paid_cents"],
                        method=PaymentMethod.BANK_TRANSFER.value,
                        reference="Imported",
                    )
                )
                payments += 1
        db.commit()
    except BaseException:
        db.rollback()
        raise
    remove_upload(path)
    return {
        "created": len(importable),
        "skipped": len(rows) - len(importable),
        "payments": payments,
        "created_categories": names.created["category"],
        "created_rooms": names.created["room"],
        "created_suppliers": names.created["supplier"],
    }


# --- Template -----------------------------------------------------------------------------


def template_workbook(currency: str = "EUR") -> Workbook:
    """Empty import sheet with one example row, plus an Instructions sheet."""
    money = excel_export.money_format(currency)
    kinds = {"amount": "money", "percent": "pct", "date": "date"}
    columns = [
        excel_export.Column(f.label, kinds.get(f.kind, "text"), 26 if f.key == "title" else 15)
        for f in FIELDS
    ]
    today = date.today()
    example = {
        "title": "Kitchen cabinets",
        "amount": 1_250_000,
        "vat_rate": 0.21,
        "category": "Kitchen",
        "subcategory": "Cabinets",
        "room": "Kitchen",
        "supplier": "Example Kitchens",
        "stage": "Ordered",
        "expense_date": today,
        "due_date": None,
        "quote_ref": "Q-2026-014",
        "invoice_no": None,
        "description": "Oak fronts, soft-close hinges",
        "notes": "Deposit paid",
        "paid_amount": 300_000,
        "paid_date": today,
    }
    wb = Workbook()
    ws = wb.active
    ws.title = "Expenses"
    excel_export.write_table(ws, columns, [[example[f.key] for f in FIELDS]], money)

    help_ws = wb.create_sheet("Instructions")
    help_columns = [
        excel_export.Column("Column", width=16),
        excel_export.Column("Required", width=10),
        excel_export.Column("What to enter", width=100),
    ]
    records = [[f.label, "Yes" if f.required else "No", f.help] for f in FIELDS]
    last = excel_export.write_table(help_ws, help_columns, records, money)
    notes = [
        "The first non-empty row holds the column names; columns may be in any order.",
        "Column names in Dutch work too (omschrijving, bedrag, btw, leverancier, ...).",
        "Rows with an error are skipped; you see them in the preview before importing.",
    ]
    for row, note in enumerate(notes, last + 2):
        help_ws.cell(row, 1, note)
    return wb
