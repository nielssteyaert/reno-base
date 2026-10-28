"""Excel exports: the full budget workbook, single-report workbooks and selected expenses.

Every figure comes from app.services.budget; this module only lays them out. Money is written
as numbers in major currency units (cents / 100) so Excel can sum, filter and chart them.
"""

from __future__ import annotations

from collections.abc import Callable, Iterable, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, Cell
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Category, Milestone, Payment, Room, Settings, Supplier
from app.services import budget
from app.services.budget import Filters, Row

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

PCT_FORMAT = "0.0%"
DATE_FORMAT = "dd/mm/yyyy"
DATETIME_FORMAT = "dd/mm/yyyy hh:mm"
MONTH_FORMAT = "mmm yyyy"

HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
OVER_FILL = PatternFill("solid", fgColor="F4B6B6")
WARNING_FILL = PatternFill("solid", fgColor="FFE08A")
BOLD = Font(bold=True)

_SYMBOLS = {"EUR": "€", "USD": "$", "GBP": "£", "JPY": "¥", "INR": "₹"}

STATUS_LABELS = {
    "estimate": "Estimate",
    "quote": "Quote",
    "ordered": "Ordered",
    "invoiced": "Invoiced",
    "paid": "Paid",
}
PAYMENT_LABELS = {"unpaid": "Unpaid", "partial": "Partial", "paid": "Paid"}
METHOD_LABELS = {"bank_transfer": "Bank transfer", "card": "Card", "cash": "Cash", "other": "Other"}
MILESTONE_LABELS = {
    "planned": "Planned",
    "in_progress": "In progress",
    "done": "Done",
    "late": "Late",
}


def money_format(currency: str) -> str:
    """Number format for money cells, e.g. `#,##0.00 "€"`; unknown currencies show their code."""
    code = (currency or "EUR").upper()
    return f'#,##0.00 "{_SYMBOLS.get(code, code)}"'


# --- Table writing ------------------------------------------------------------------------


@dataclass(frozen=True)
class Column:
    header: str
    kind: str = "text"  # text, int, money (cents), pct (ratio), date, datetime, month
    width: int = 14


def put(cell: Cell, value: Any, kind: str, money: str) -> None:
    """Write one value with the number format that belongs to its kind."""
    if value is None:
        return
    if kind == "money":
        cell.value = value / 100
        cell.number_format = money
    elif kind == "pct":
        cell.value = value
        cell.number_format = PCT_FORMAT
    elif kind in ("date", "datetime", "month"):
        cell.value = value
        cell.number_format = {
            "date": DATE_FORMAT,
            "datetime": DATETIME_FORMAT,
            "month": MONTH_FORMAT,
        }[kind]
    elif isinstance(value, str):
        cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)
        # openpyxl stores strings starting with "=" as formulas; user text must stay text.
        cell.data_type = "s"
    else:
        cell.value = value


def write_header(ws: Worksheet, columns: Sequence[Column], row: int = 1) -> None:
    for col, column in enumerate(columns, 1):
        cell = ws.cell(row, col, column.header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        ws.column_dimensions[get_column_letter(col)].width = column.width


def write_table(
    ws: Worksheet,
    columns: Sequence[Column],
    records: Iterable[Sequence[Any]],
    money: str,
    *,
    top: int = 1,
    finish: bool = True,
) -> int:
    """Header at `top` and one row per record below it; returns the last row written.

    With `finish`, the header is frozen and the data range gets an autofilter.
    """
    write_header(ws, columns, top)
    last = top
    for last, record in enumerate(records, top + 1):
        for col, (column, value) in enumerate(zip(columns, record, strict=True), 1):
            put(ws.cell(last, col), value, column.kind, money)
    if finish:
        ws.freeze_panes = ws.cell(top + 1, 1)
        ws.auto_filter.ref = f"A{top}:{get_column_letter(len(columns))}{last}"
    return last


def _fill_row(ws: Worksheet, row: int, ncols: int, fill: PatternFill | None) -> None:
    if fill is None:
        return
    for col in range(1, ncols + 1):
        ws.cell(row, col).fill = fill


def _health_fill(line: dict) -> PatternFill | None:
    return {"over": OVER_FILL, "warning": WARNING_FILL}.get(line["health"])


def _budget_chart(
    ws: Worksheet,
    title: str,
    *,
    labels_col: int,
    budget_col: int,
    forecast_col: int,
    header_row: int,
    last_row: int,
    anchor: str,
    money: str,
) -> None:
    """Clustered columns of budget vs forecast; series titles come from the header row."""
    chart = BarChart()
    chart.type = "col"
    chart.grouping = "clustered"
    chart.title = title
    chart.y_axis.number_format = money.replace(".00", "")
    for col in (budget_col, forecast_col):
        chart.add_data(
            Reference(ws, min_col=col, min_row=header_row, max_row=last_row),
            titles_from_data=True,
        )
    chart.set_categories(
        Reference(ws, min_col=labels_col, min_row=header_row + 1, max_row=last_row)
    )
    _show_axes(chart)
    chart.width, chart.height = 20, 10
    ws.add_chart(chart, anchor)


def _show_axes(chart: BarChart | LineChart) -> None:
    # openpyxl 3.1 marks axes as deleted by default, which hides them in Excel.
    chart.x_axis.delete = False
    chart.y_axis.delete = False


# --- Sheets -------------------------------------------------------------------------------


@dataclass
class _Context:
    db: Session
    settings: Settings
    rows: list[Row]
    filters: Filters
    money: str


SheetBuilder = Callable[[Workbook, _Context], None]

BUDGET_COLUMNS = [
    Column("Budget", "money"),
    Column("Estimated", "money"),
    Column("Quoted", "money"),
    Column("Committed", "money"),
    Column("Paid", "money"),
    Column("Outstanding", "money"),
    Column("Forecast", "money"),
    Column("Remaining", "money"),
    Column("% committed", "pct", 12),
]

EXPENSE_COLUMNS = [
    Column("ID", "int", 7),
    Column("Title", width=32),
    Column("Category", width=18),
    Column("Subcategory", width=18),
    Column("Room", width=16),
    Column("Supplier", width=22),
    Column("Milestone", width=18),
    Column("Status", width=11),
    Column("Payment status", width=15),
    Column("Date", "date", 12),
    Column("Due date", "date", 12),
    Column("Excl. VAT", "money"),
    Column("VAT %", "pct", 8),
    Column("VAT", "money", 12),
    Column("Incl. VAT", "money"),
    Column("Paid", "money"),
    Column("Outstanding", "money"),
    Column("Quote ref", width=14),
    Column("Invoice no", width=14),
    Column("Description", width=40),
    Column("Notes", width=40),
]
_EXPENSE_TOTALS = {"Excl. VAT", "VAT", "Incl. VAT", "Paid", "Outstanding"}
_FORECAST_INDEX = [c.header for c in BUDGET_COLUMNS].index("Forecast")


def _budget_values(line: dict) -> list:
    return [
        line["budget"],
        line["estimated"],
        line["quoted"],
        line["committed"],
        line["paid"],
        line["outstanding"],
        line["forecast"],
        line["remaining"],
        line["pct_committed"],
    ]


def describe_filters(db: Session, filters: Filters) -> list[str]:
    """Human-readable list of the active report filters."""
    parts = []
    if filters.date_from:
        parts.append(f"From {filters.date_from:%d/%m/%Y}")
    if filters.date_to:
        parts.append(f"To {filters.date_to:%d/%m/%Y}")
    named = (
        ("Category", Category, filters.category_ids),
        ("Room", Room, filters.room_ids),
        ("Supplier", Supplier, filters.supplier_ids),
        ("Milestone", Milestone, filters.milestone_ids),
    )
    for label, model, ids in named:
        if ids:
            names = [getattr(db.get(model, i), "name", None) or f"#{i}" for i in ids]
            parts.append(f"{label}: {', '.join(names)}")
    if filters.statuses:
        parts.append("Status: " + ", ".join(STATUS_LABELS.get(s, s) for s in filters.statuses))
    if filters.q:
        parts.append(f'Search: "{filters.q}"')
    return parts


def _summary_sheet(wb: Workbook, ctx: _Context) -> None:
    ws = wb.create_sheet("Summary")
    summary = budget.project_summary(ctx.db, ctx.rows)
    ws["A1"] = ctx.settings.project_name
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Exported"
    put(ws["B2"], datetime.now().replace(microsecond=0), "datetime", ctx.money)
    row = 3
    for i, text in enumerate(describe_filters(ctx.db, ctx.filters)):
        ws.cell(row, 1, "Filters" if i == 0 else None)
        put(ws.cell(row, 2), text, "text", ctx.money)
        row += 1

    kpis = [
        ("Total budget", ctx.settings.total_budget_cents, "money"),
        ("Contingency %", summary["contingency_pct"] / 100, "pct"),
        ("Contingency reserve", summary["contingency_reserve"], "money"),
        ("Working budget", summary["working_budget"], "money"),
        ("Estimated", summary["estimated"], "money"),
        ("Quoted", summary["quoted"], "money"),
        ("Committed", summary["committed"], "money"),
        ("Paid", summary["paid"], "money"),
        ("Outstanding", summary["outstanding"], "money"),
        ("Forecast", summary["forecast"], "money"),
        ("Remaining", summary["remaining"], "money"),
        ("% committed", summary["pct_committed"], "pct"),
    ]
    kpi_top = row + 1
    write_header(ws, [Column("Metric", width=26), Column("Value", width=20)], kpi_top)
    for row, (label, value, kind) in enumerate(kpis, kpi_top + 1):
        ws.cell(row, 1, label)
        put(ws.cell(row, 2), value, kind, ctx.money)
    ws.freeze_panes = ws.cell(kpi_top + 1, 1)
    ws.auto_filter.ref = f"A{kpi_top}:B{row}"

    status_columns = [
        Column("Status", width=26),
        Column("Count", "int", 20),
        Column("Amount", "money"),
    ]
    funnel = [
        (STATUS_LABELS[item["status"]], item["count"], item["amount"]) for item in summary["funnel"]
    ]
    write_table(ws, status_columns, funnel, ctx.money, top=row + 2, finish=False)


def _expense_record(row: Row) -> list:
    expense, figures = row.expense, row.figures
    category = expense.category
    if category is not None and category.parent is not None:
        top_name, sub_name = category.parent.name, category.name
    else:
        top_name, sub_name = (category.name if category else None), None
    return [
        expense.id,
        expense.title,
        top_name,
        sub_name,
        expense.room.name if expense.room else None,
        expense.supplier.name if expense.supplier else None,
        expense.milestone.name if expense.milestone else None,
        STATUS_LABELS.get(figures.status, figures.status),
        PAYMENT_LABELS.get(figures.payment_status, figures.payment_status),
        expense.expense_date,
        expense.due_date,
        figures.excl,
        expense.vat_rate / 100,
        figures.vat,
        figures.incl,
        figures.paid,
        figures.outstanding,
        expense.quote_ref,
        expense.invoice_no,
        expense.description,
        expense.notes,
    ]


def _expenses_sheet(wb: Workbook, ctx: _Context) -> None:
    ws = wb.create_sheet("Expenses")
    last = write_table(ws, EXPENSE_COLUMNS, map(_expense_record, ctx.rows), ctx.money)
    if not ctx.rows:
        return
    # SUBTOTAL(109, ...) sums only the rows left visible by the autofilter.
    total_row = last + 1
    ws.cell(total_row, 2, "Total").font = BOLD
    for col, column in enumerate(EXPENSE_COLUMNS, 1):
        if column.header in _EXPENSE_TOTALS:
            letter = get_column_letter(col)
            cell = ws.cell(total_row, col, f"=SUBTOTAL(109,{letter}2:{letter}{last})")
            cell.number_format = ctx.money
            cell.font = BOLD


def _payments_sheet(wb: Workbook, ctx: _Context, *, chronological: bool) -> None:
    """Payments of the exported expenses: by date, or grouped in the order of the expenses."""
    ws = wb.create_sheet("Payments")
    position = {row.expense.id: i for i, row in enumerate(ctx.rows)}
    expenses = {row.expense.id: row.expense for row in ctx.rows}
    payments = ctx.db.scalars(select(Payment).where(Payment.expense_id.in_(list(position)))).all()
    if chronological:
        payments = sorted(payments, key=lambda p: (p.date, p.id))
    else:
        payments = sorted(payments, key=lambda p: (position[p.expense_id], p.date, p.id))
    columns = [
        Column("Date", "date", 12),
        Column("Expense ID", "int", 11),
        Column("Expense", width=32),
        Column("Supplier", width=22),
        Column("Amount", "money"),
        Column("Method", width=14),
        Column("Reference", width=18),
        Column("Note", width=40),
    ]
    records = []
    for payment in payments:
        expense = expenses[payment.expense_id]
        records.append(
            [
                payment.date,
                expense.id,
                expense.title,
                expense.supplier.name if expense.supplier else None,
                payment.amount_cents,
                METHOD_LABELS.get(payment.method, payment.method),
                payment.reference,
                payment.note,
            ]
        )
    write_table(ws, columns, records, ctx.money)


def _categories_sheet(wb: Workbook, ctx: _Context) -> None:
    ws = wb.create_sheet("Categories")
    tree = budget.by_category(ctx.db, ctx.rows)
    columns = [Column("Category", width=22), Column("Subcategory", width=22), *BUDGET_COLUMNS]
    records, styles = [], []
    for parent in tree:
        records.append([parent["name"], None, *_budget_values(parent)])
        styles.append((parent, True))
        for child in parent.get("children", []):
            records.append([parent["name"], child["name"], *_budget_values(child)])
            styles.append((child, False))
    last = write_table(ws, columns, records, ctx.money)
    for row, (line, is_parent) in enumerate(styles, 2):
        _fill_row(ws, row, len(columns), _health_fill(line))
        if is_parent:
            ws.cell(row, 1).font = BOLD
        else:
            ws.cell(row, 2).alignment = Alignment(indent=1)
    if not tree:
        return

    # Chart source: top-level categories only, below the table so the autofilter never hides it.
    chart_top = last + 3
    ws.cell(chart_top - 1, 1, "Chart data: top-level categories").font = BOLD
    chart_columns = [
        Column("Category", width=22),
        Column("Budget", "money"),
        Column("Forecast", "money"),
    ]
    chart_records = [[line["name"], line["budget"], line["forecast"]] for line in tree]
    for col, column in enumerate(chart_columns, 1):
        ws.cell(chart_top, col, column.header).font = HEADER_FONT
        ws.cell(chart_top, col).fill = HEADER_FILL
    for row, record in enumerate(chart_records, chart_top + 1):
        for col, (column, value) in enumerate(zip(chart_columns, record, strict=True), 1):
            put(ws.cell(row, col), value, column.kind, ctx.money)
    _budget_chart(
        ws,
        "Budget vs forecast by category",
        labels_col=1,
        budget_col=2,
        forecast_col=3,
        header_row=chart_top,
        last_row=chart_top + len(chart_records),
        anchor=f"{get_column_letter(len(columns) + 2)}2",
        money=ctx.money,
    )


def _rooms_sheet(wb: Workbook, ctx: _Context) -> None:
    ws = wb.create_sheet("Rooms")
    lines = budget.by_room(ctx.db, ctx.rows)
    columns = [Column("Room", width=22), *BUDGET_COLUMNS]
    records = [[line["name"], *_budget_values(line)] for line in lines]
    last = write_table(ws, columns, records, ctx.money)
    for row, line in enumerate(lines, 2):
        _fill_row(ws, row, len(columns), _health_fill(line))
    if not lines:
        return
    _budget_chart(
        ws,
        "Budget vs forecast by room",
        labels_col=1,
        budget_col=2,
        forecast_col=2 + _FORECAST_INDEX,
        header_row=1,
        last_row=last,
        anchor=f"{get_column_letter(len(columns) + 2)}2",
        money=ctx.money,
    )


def _suppliers_sheet(wb: Workbook, ctx: _Context) -> None:
    ws = wb.create_sheet("Suppliers")
    suppliers = {s.id: s for s in ctx.db.scalars(select(Supplier)).all()}
    columns = [
        Column("Supplier", width=26),
        Column("Contact", width=20),
        Column("Email", width=26),
        Column("Phone", width=16),
        Column("Expenses", "int", 10),
        Column("Quoted", "money"),
        Column("Committed", "money"),
        Column("Paid", "money"),
        Column("Outstanding", "money"),
        Column("Last payment", "date", 13),
    ]
    records = []
    for line in budget.by_supplier(ctx.db, ctx.rows):
        supplier = suppliers.get(line["id"])
        records.append(
            [
                line["name"],
                supplier.contact_person if supplier else None,
                supplier.email if supplier else None,
                supplier.phone if supplier else None,
                line["count"],
                line["quoted"],
                line["committed"],
                line["paid"],
                line["outstanding"],
                line.get("last_payment"),
            ]
        )
    write_table(ws, columns, records, ctx.money)


def _milestones_sheet(wb: Workbook, ctx: _Context) -> None:
    ws = wb.create_sheet("Milestones")
    milestones = ctx.db.scalars(
        select(Milestone).order_by(
            Milestone.sort_order,
            Milestone.planned_start.is_(None),
            Milestone.planned_start,
            Milestone.id,
        )
    ).all()
    rooms = {r.id: r.name for r in ctx.db.scalars(select(Room)).all()}
    linked = budget.by_milestone(ctx.db, ctx.rows)
    today = date.today()
    columns = [
        Column("Name", width=28),
        Column("Room", width=18),
        Column("Status", width=12),
        Column("Planned start", "date", 13),
        Column("Planned end", "date", 13),
        Column("Actual start", "date", 13),
        Column("Actual end", "date", 13),
        Column("Linked forecast", "money", 16),
        Column("Linked paid", "money", 14),
    ]
    records = []
    for milestone in milestones:
        line = linked.get(milestone.id, {})
        status = budget.milestone_status(milestone, today)
        records.append(
            [
                milestone.name,
                rooms.get(milestone.room_id),
                MILESTONE_LABELS.get(status, status),
                milestone.planned_start,
                milestone.planned_end,
                milestone.actual_start,
                milestone.actual_end,
                line.get("forecast", 0),
                line.get("paid", 0),
            ]
        )
    write_table(ws, columns, records, ctx.money)


def _monthly_sheet(wb: Workbook, ctx: _Context) -> None:
    ws = wb.create_sheet("Monthly")
    months = budget.monthly(ctx.db, ctx.rows)
    total_budget = ctx.settings.total_budget_cents or 0
    columns = [
        Column("Month", "month", 12),
        Column("Paid", "money"),
        Column("Committed", "money"),
        Column("Cumulative paid", "money", 17),
        Column("Cumulative committed", "money", 21),
        Column("Total budget", "money"),
    ]
    records = [
        [
            date(*map(int, m["month"].split("-")), 1),
            m["paid"],
            m["committed"],
            m["cumulative_paid"],
            m["cumulative_committed"],
            total_budget,
        ]
        for m in months
    ]
    last = write_table(ws, columns, records, ctx.money)
    if not records:
        return
    chart = LineChart()
    chart.title = "Cumulative spending"
    chart.y_axis.number_format = ctx.money.replace(".00", "")
    chart.x_axis.number_format = MONTH_FORMAT
    series_cols = [4, 5] + ([6] if total_budget else [])
    for col in series_cols:
        chart.add_data(Reference(ws, min_col=col, min_row=1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
    _show_axes(chart)
    chart.width, chart.height = 22, 10
    ws.add_chart(chart, f"{get_column_letter(len(columns) + 2)}2")


# --- Public entry points ------------------------------------------------------------------

REPORTS: dict[str, tuple[SheetBuilder, ...]] = {
    "categories": (_categories_sheet,),
    "rooms": (_rooms_sheet,),
    "suppliers": (_suppliers_sheet,),
    "monthly": (_monthly_sheet,),
    "forecast": (_summary_sheet, _categories_sheet),
}


def _context(db: Session, rows: list[Row], filters: Filters) -> _Context:
    settings = budget.get_settings(db)
    return _Context(db, settings, rows, filters, money_format(settings.currency))


def _build(ctx: _Context, builders: Iterable[SheetBuilder]) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    for build in builders:
        build(wb, ctx)
    return wb


def full_workbook(db: Session, filters: Filters) -> Workbook:
    """Every sheet: summary, expenses, payments and all reports, for the filtered expenses."""
    ctx = _context(db, budget.load_rows(db, filters), filters)
    return _build(
        ctx,
        (
            _summary_sheet,
            _expenses_sheet,
            lambda wb, c: _payments_sheet(wb, c, chronological=True),
            _categories_sheet,
            _rooms_sheet,
            _suppliers_sheet,
            _milestones_sheet,
            _monthly_sheet,
        ),
    )


def report_workbook(db: Session, name: str, filters: Filters) -> Workbook:
    """Workbook with one report's sheet(s). Raises KeyError for an unknown report name."""
    builders = REPORTS[name]
    return _build(_context(db, budget.load_rows(db, filters), filters), builders)


def expenses_workbook(db: Session, ids: Sequence[int]) -> Workbook:
    """Expenses and Payments sheets for exactly these expenses, in the given order."""
    by_id = {row.expense.id: row for row in budget.load_rows(db)}
    rows = [by_id[i] for i in dict.fromkeys(ids) if i in by_id]
    ctx = _context(db, rows, Filters())
    return _build(ctx, (_expenses_sheet, lambda wb, c: _payments_sheet(wb, c, chronological=False)))


def to_stream(wb: Workbook) -> BytesIO:
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
