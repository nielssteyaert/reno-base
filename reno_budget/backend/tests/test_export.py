from __future__ import annotations

import io
from datetime import date

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.services.excel_export import XLSX_MEDIA_TYPE

SHEETS = [
    "Summary",
    "Expenses",
    "Payments",
    "Categories",
    "Rooms",
    "Suppliers",
    "Milestones",
    "Monthly",
]


def _expense(client: TestClient, **fields) -> dict:
    body = {"title": "Item", "amount_cents": 100_000, "vat_rate": 21, **fields}
    response = client.post("/api/expenses", json=body)
    assert response.status_code == 201, response.text
    return response.json()


def _workbook(response) -> Workbook:
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE
    return load_workbook(io.BytesIO(response.content))


def _records(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    return [dict(zip(rows[0], row, strict=True)) for row in rows[1:]]


@pytest.fixture
def project(client: TestClient, seeded: dict) -> dict:
    oven = _expense(
        client,
        title="Oven",
        amount_cents=123_456,
        vat_rate=6,
        stage="ordered",
        category_id=seeded["appliances"]["id"],
        room_id=seeded["room"]["id"],
        supplier_id=seeded["supplier"]["id"],
        expense_date="2026-02-10",
    )
    wiring = _expense(
        client,
        title="Wiring",
        amount_cents=1_100_000,
        stage="invoiced",
        category_id=seeded["electrical"]["id"],
        expense_date="2026-03-05",
        due_date="2026-04-01",
    )
    tiles = _expense(client, title="Tiles", stage="quote", amount_cents=50_000)
    client.post(
        f"/api/expenses/{oven['id']}/payments", json={"date": "2026-02-20", "amount_cents": 50_000}
    )
    client.post(
        f"/api/expenses/{wiring['id']}/payments",
        json={"date": "2026-03-10", "amount_cents": 1_331_000, "reference": "INV-1"},
    )
    return {"oven": oven, "wiring": wiring, "tiles": tiles, **seeded}


def test_workbook_sheets_and_formatting(client: TestClient, project: dict):
    response = client.get("/api/export/workbook")
    wb = _workbook(response)
    today = date.today().isoformat()
    assert response.headers["content-disposition"] == (
        f'attachment; filename="renovation-budget-{today}.xlsx"'
    )
    assert wb.sheetnames == SHEETS

    expenses = wb["Expenses"]
    assert expenses["A1"].font.bold
    assert expenses.freeze_panes == "A2"
    assert expenses.auto_filter.ref == "A1:U4"
    assert expenses["O2"].number_format == '#,##0.00 "€"'
    assert expenses["M2"].number_format == "0.0%"
    assert expenses["J3"].number_format == "dd/mm/yyyy"  # row 2 (Tiles) has no date
    for ws in wb:
        assert ws.freeze_panes is not None, ws.title
        assert ws.auto_filter.ref, ws.title

    summary: dict = {}
    for label, value, _ in wb["Summary"].iter_rows(values_only=True):
        summary.setdefault(label, value)  # KPI table first; the status table repeats "Paid"
    assert summary["Total budget"] == 100_000
    assert summary["Contingency %"] == pytest.approx(0.1)
    api = client.get("/api/reports/summary").json()
    assert summary["Forecast"] == pytest.approx(api["forecast"] / 100)
    assert summary["Paid"] == pytest.approx(api["paid"] / 100)


def test_expense_values_match_api(client: TestClient, project: dict):
    wb = _workbook(client.get("/api/export/workbook"))
    ws = wb["Expenses"]
    rows = list(ws.iter_rows(values_only=True))
    header, data, totals = rows[0], rows[1:-1], rows[-1]
    records = {row[0]: dict(zip(header, row, strict=True)) for row in data}
    api = client.get("/api/expenses").json()
    assert sorted(records) == sorted(e["id"] for e in api)
    for expense in api:
        record = records[expense["id"]]
        assert record["Title"] == expense["title"]
        assert record["Excl. VAT"] == pytest.approx(expense["amount_excl_cents"] / 100)
        assert record["VAT"] == pytest.approx(expense["vat_cents"] / 100)
        assert record["Incl. VAT"] == pytest.approx(expense["amount_incl_cents"] / 100)
        assert record["Paid"] == pytest.approx(expense["paid_cents"] / 100)
        assert record["Outstanding"] == pytest.approx(expense["outstanding_cents"] / 100)
        assert record["VAT %"] == pytest.approx(expense["vat_rate"] / 100)

    oven = records[project["oven"]["id"]]
    assert (oven["Category"], oven["Subcategory"]) == ("Kitchen", "Appliances")
    assert oven["Room"] == "Ground floor"
    assert oven["Supplier"] == "ACME Builders"
    assert oven["Date"].date() == date(2026, 2, 10)
    assert records[project["wiring"]["id"]]["Status"] == "Paid"

    assert totals[1] == "Total"
    assert totals[header.index("Incl. VAT")] == "=SUBTOTAL(109,O2:O4)"
    assert totals[header.index("Paid")] == "=SUBTOTAL(109,P2:P4)"

    payments = _records(wb["Payments"])
    assert [p["Expense ID"] for p in payments] == [project["oven"]["id"], project["wiring"]["id"]]
    assert payments[1]["Amount"] == pytest.approx(13_310)
    assert payments[1]["Reference"] == "INV-1"


def test_workbook_respects_filters(client: TestClient, project: dict):
    wb = _workbook(
        client.get("/api/export/workbook", params={"category_id": project["kitchen"]["id"]})
    )
    assert [r["Title"] for r in _records(wb["Expenses"]) if r["ID"]] == ["Oven"]
    summary_text = [row[1] for row in wb["Summary"].iter_rows(values_only=True)]
    assert "Category: Kitchen" in summary_text


def test_selected_expenses_keep_order(client: TestClient, project: dict):
    ids = [project["wiring"]["id"], project["tiles"]["id"], project["oven"]["id"]]
    response = client.post("/api/export/expenses", json={"ids": ids})
    wb = _workbook(response)
    assert "renovation-expenses-" in response.headers["content-disposition"]
    assert wb.sheetnames == ["Expenses", "Payments"]
    assert [r["ID"] for r in _records(wb["Expenses"])][:-1] == ids
    payments = _records(wb["Payments"])
    assert [p["Expense ID"] for p in payments] == [project["wiring"]["id"], project["oven"]["id"]]

    wb = _workbook(client.post("/api/export/expenses", json={"ids": [project["tiles"]["id"]]}))
    assert [r["Title"] for r in _records(wb["Expenses"])][:-1] == ["Tiles"]
    assert _records(wb["Payments"]) == []


def test_single_reports(client: TestClient, project: dict):
    response = client.get("/api/export/report/categories")
    wb = _workbook(response)
    assert wb.sheetnames == ["Categories"]
    today = date.today().isoformat()
    assert (
        f'filename="renovation-categories-{today}.xlsx"' in response.headers["content-disposition"]
    )
    assert _workbook(client.get("/api/export/report/forecast")).sheetnames == [
        "Summary",
        "Categories",
    ]
    for name, sheet in (("rooms", "Rooms"), ("suppliers", "Suppliers"), ("monthly", "Monthly")):
        assert _workbook(client.get(f"/api/export/report/{name}")).sheetnames == [sheet]


def test_unknown_report_is_404(client: TestClient):
    assert client.get("/api/export/report/nope").status_code == 404


def test_categories_sheet_rows_fills_and_chart(client: TestClient, project: dict):
    # Electrical: 10k budget, 11k + 21% VAT invoiced -> over budget.
    wb = load_workbook(io.BytesIO(client.get("/api/export/report/categories").content))
    ws = wb["Categories"]
    assert len(ws._charts) == 1
    rows = list(ws.iter_rows(min_row=1, max_row=4, values_only=True))
    assert rows[0][:3] == ("Category", "Subcategory", "Budget")
    labels = [(r[0], r[1]) for r in rows[1:]]
    assert labels == [("Electrical", None), ("Kitchen", None), ("Kitchen", "Appliances")]
    assert ws["A2"].fill.fgColor.rgb.endswith("F4B6B6")
    assert ws["A3"].fill.fill_type is None
    assert rows[1][2] == 10_000
    assert rows[1][8] == pytest.approx(13_310)

    rooms = _workbook(client.get("/api/export/report/rooms"))["Rooms"]
    assert len(rooms._charts) == 1
    monthly = _workbook(client.get("/api/export/report/monthly"))["Monthly"]
    assert len(monthly._charts) == 1
    assert monthly["A2"].value.date() == date(2026, 2, 1)


def test_text_is_never_a_formula(client: TestClient):
    _expense(client, title='=HYPERLINK("http://example.com")', notes="+1")
    wb = _workbook(client.get("/api/export/workbook"))
    cell = wb["Expenses"]["B2"]
    assert cell.data_type == "s"
    assert cell.value.startswith("=HYPERLINK")


def test_currency_symbol_follows_settings(client: TestClient):
    client.patch("/api/settings", json={"currency": "CHF"})
    _expense(client)
    wb = _workbook(client.get("/api/export/workbook"))
    assert wb["Expenses"]["O2"].number_format == '#,##0.00 "CHF"'
