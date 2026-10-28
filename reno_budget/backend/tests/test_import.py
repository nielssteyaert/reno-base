from __future__ import annotations

import io
import os
from datetime import date, datetime
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.models import Stage
from app.services.importer import (
    guess_mapping,
    parse_amount,
    parse_date,
    parse_stage,
    parse_vat,
    split_category,
)

DUTCH_HEADERS = [
    "Omschrijving",
    "Bedrag",
    "BTW %",
    "Categorie",
    "Ruimte",
    "Leverancier",
    "Status",
    "Datum",
    "Vervaldatum",
    "Offerte",
    "Factuur",
    "Opmerking",
    "Betaald",
    "Betaaldatum",
]


def _xlsx(rows: list[list], sheets: dict[str, list[list]] | None = None) -> bytes:
    wb = Workbook()
    wb.active.title = "Blad1"
    for row in rows:
        wb.active.append(row)
    for name, extra in (sheets or {}).items():
        ws = wb.create_sheet(name)
        for row in extra:
            ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _upload(client: TestClient, content: bytes, filename: str = "budget.xlsx") -> dict:
    response = client.post("/api/import/upload", files={"file": (filename, content)})
    assert response.status_code == 200, response.text
    return response.json()


# --- Parsing ------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("raw", "cents"),
    [
        (1234.5, 123_450),
        (12, 1_200),
        ("1.234,56", 123_456),
        ("1,234.56", 123_456),
        ("€ 1 234,50", 123_450),
        ("1 234,50 EUR", 123_450),
        ("1234", 123_400),
        ("1.234", 123_400),
        ("1,5", 150),
        ("12.345.678,9", 1_234_567_890),
        ("12,3", 1_230),
        ("-", None),
        ("", None),
        (None, None),
    ],
)
def test_parse_amount(raw, cents):
    assert parse_amount(raw) == cents


@pytest.mark.parametrize("raw", ["abc", "-12", "(5,00)", "12%", -3, "1,2,3x"])
def test_parse_amount_rejects(raw):
    with pytest.raises(ValueError):
        parse_amount(raw)


@pytest.mark.parametrize(
    ("raw", "rate"),
    [
        (21, 21),
        ("21%", 21),
        ("21 %", 21),
        (0.21, 21),
        ("0,06", 6),
        ("6,0", 6),
        (0, 0),
        (None, None),
    ],
)
def test_parse_vat(raw, rate):
    assert parse_vat(raw) == rate


@pytest.mark.parametrize("raw", ["150", "abc", "-6"])
def test_parse_vat_rejects(raw):
    with pytest.raises(ValueError):
        parse_vat(raw)


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        (datetime(2026, 3, 4, 0, 0), date(2026, 3, 4)),
        (date(2026, 3, 4), date(2026, 3, 4)),
        ("04/03/2026", date(2026, 3, 4)),
        ("4-3-2026", date(2026, 3, 4)),
        ("04.03.26", date(2026, 3, 4)),
        ("2026-03-04", date(2026, 3, 4)),
        ("2026-03-04 00:00:00", date(2026, 3, 4)),
        (46085, date(2026, 3, 4)),
        (None, None),
    ],
)
def test_parse_date(raw, expected):
    assert parse_date(raw) == expected


@pytest.mark.parametrize("raw", ["31/02/2026", "next week", "2026/13/01", "13/13/26"])
def test_parse_date_rejects(raw):
    with pytest.raises(ValueError):
        parse_date(raw)


def test_parse_stage_and_category():
    assert parse_stage("Offerte") == (Stage.QUOTE, False)
    assert parse_stage("BESTELD") == (Stage.ORDERED, False)
    assert parse_stage("gefactureerd") == (Stage.INVOICED, False)
    assert parse_stage("Raming") == (Stage.ESTIMATE, False)
    assert parse_stage("Betaald") == (Stage.INVOICED, True)
    assert parse_stage(None) == (None, False)
    with pytest.raises(ValueError):
        parse_stage("misschien")
    assert split_category("Keuken > Toestellen") == ("Keuken", "Toestellen")
    assert split_category("Keuken / Toestellen") == ("Keuken", "Toestellen")
    assert split_category("Keuken") == ("Keuken", None)


def test_guess_mapping_dutch_headers():
    mapping = guess_mapping(DUTCH_HEADERS)
    assert mapping == {
        "title": 0,
        "amount": 1,
        "vat_rate": 2,
        "category": 3,
        "subcategory": None,
        "room": 4,
        "supplier": 5,
        "stage": 6,
        "expense_date": 7,
        "due_date": 8,
        "quote_ref": 9,
        "invoice_no": 10,
        "description": None,
        "notes": 11,
        "paid_amount": 12,
        "paid_date": 13,
    }


def test_guess_mapping_title_and_description():
    mapping = guess_mapping(["Titel", "Omschrijving", "Prijs (incl. BTW)", "Aannemer", "KAMER"])
    assert mapping["title"] == 0
    assert mapping["description"] == 1
    assert mapping["amount"] == 2
    assert mapping["supplier"] == 3
    assert mapping["room"] == 4
    assert mapping["vat_rate"] is None


# --- Upload and inspect -------------------------------------------------------------------


def test_upload_inspects_first_non_empty_row(client: TestClient, data_dir: Path):
    rows = [[], [None, None], DUTCH_HEADERS] + [
        [f"Item {i}", i * 10, 21, None, None, None, None, None, None] for i in range(12)
    ]
    other = [["Title", "Amount"], ["Paint", "12,50"]]
    result = _upload(client, _xlsx(rows, {"Andere": other}))
    assert result["filename"] == "budget.xlsx"
    assert result["sheets"] == ["Blad1", "Andere"]
    assert result["sheet"] == "Blad1"
    assert result["headers"] == DUTCH_HEADERS
    assert result["row_count"] == 12
    assert len(result["sample"]) == 10
    assert result["sample"][1][:3] == ["Item 1", "10", "21"]
    assert result["sample"][1][3] == ""
    assert result["mapping"]["title"] == 0
    stored = list((data_dir / "imports").glob(f"{result['token']}.*"))
    assert {p.suffix for p in stored} == {".xlsx", ".name"}

    other_sheet = client.post(
        "/api/import/inspect", json={"token": result["token"], "sheet": "Andere"}
    ).json()
    assert other_sheet["sheet"] == "Andere"
    assert other_sheet["headers"] == ["Title", "Amount"]
    assert other_sheet["sample"] == [["Paint", "12,50"]]
    assert other_sheet["filename"] == "budget.xlsx"


def test_upload_rejects_other_files_and_removes_stale(client: TestClient, data_dir: Path):
    response = client.post("/api/import/upload", files={"file": ("notes.txt", b"hello")})
    assert response.status_code == 415

    imports = data_dir / "imports"
    imports.mkdir(exist_ok=True)
    stale = imports / ("0" * 32 + ".csv")
    stale.write_text("Title;Amount\n")
    old = datetime(2020, 1, 1).timestamp()
    os.utime(stale, (old, old))
    _upload(client, b"Title;Amount\nX;1\n", "x.csv")
    assert not stale.exists()


def test_unknown_token_is_404(client: TestClient):
    body = {"token": "f" * 32, "mapping": {"title": 0, "amount": 1}}
    assert client.post("/api/import/preview", json=body).status_code == 404
    assert client.post("/api/import/commit", json=body).status_code == 404
    assert client.post("/api/import/inspect", json={"token": "../../etc"}).status_code == 404


def test_missing_required_mapping_is_422(client: TestClient):
    token = _upload(client, b"Title;Amount\nX;1\n", "x.csv")["token"]
    response = client.post("/api/import/preview", json={"token": token, "mapping": {"title": 0}})
    assert response.status_code == 422
    assert "Amount" in response.json()["detail"]


# --- Preview and commit -------------------------------------------------------------------


def _dutch_upload(client: TestClient, seeded: dict) -> dict:
    rows = [
        DUTCH_HEADERS,
        # 0: existing category (other case), new subcategory, existing supplier
        [
            "Oven",
            "1.000,00",
            "21%",
            "kitchen > Ovens",
            "Keuken",
            "acme builders",
            "Besteld",
            "01/02/2026",
            None,
            "OF-1",
            None,
            None,
            None,
            None,
        ],
        # 1: new category with new child, paid in full via the status word
        [
            "Tegels",
            500,
            0.06,
            "Vloeren / Tegels",
            "Badkamer",
            "Tegelhuis",
            "betaald",
            "15/02/2026",
            None,
            None,
            "F-22",
            "met plinten",
            None,
            None,
        ],
        # 2: invalid amount -> error
        ["Kapot", "veel", 21, "Vloeren"] + [None] * 10,
        # 3: unknown status -> warning, default stage; partial payment
        [
            "Kraan",
            "200",
            None,
            "Sanitair",
            "badkamer",
            "Tegelhuis",
            "misschien",
            datetime(2026, 3, 1),
            "2026-04-01",
            None,
            None,
            None,
            "50",
            "02/03/2026",
        ],
        # 4: bad date -> error
        ["Lamp", "80", 21, None, None, None, None, "31/02/2026"] + [None] * 6,
    ]
    return _upload(client, _xlsx(rows))


def test_preview_reports_unmatched_and_errors(client: TestClient, seeded: dict):
    upload = _dutch_upload(client, seeded)
    body = {"token": upload["token"], "mapping": upload["mapping"], "default_vat_rate": 21}
    result = client.post("/api/import/preview", json=body).json()

    assert result["valid_count"] == 3
    assert result["error_count"] == 2
    rows = result["rows"]
    assert [r["index"] for r in rows] == [0, 1, 2, 3, 4]
    assert rows[2]["errors"] and "veel" in rows[2]["errors"][0]
    assert rows[4]["errors"] and "31/02/2026" in rows[4]["errors"][0]
    assert rows[3]["warnings"] and "misschien" in rows[3]["warnings"][0]

    oven = rows[0]["values"]
    assert oven["amount_excl_cents"] == 100_000
    assert oven["vat_cents"] == 21_000
    assert oven["category"] == "kitchen"
    assert oven["subcategory"] == "Ovens"
    assert oven["stage"] == "ordered"
    assert oven["expense_date"] == "2026-02-01"
    tiles = rows[1]["values"]
    assert tiles["vat_rate"] == 6
    assert tiles["stage"] == "invoiced"
    assert tiles["paid_cents"] == tiles["amount_incl_cents"] == 53_000
    assert tiles["paid_date"] == "2026-02-15"
    tap = rows[3]["values"]
    assert tap["stage"] == "invoiced"
    assert tap["vat_rate"] == 21
    assert tap["paid_cents"] == 5_000

    assert result["unmatched"] == {
        "category": ["Vloeren", "Sanitair"],
        "room": ["Keuken", "Badkamer"],
        "supplier": ["Tegelhuis"],
    }
    assert result["total_incl_cents"] == 121_000 + 53_000 + 24_200

    body["skip_rows"] = [3]
    body["value_map"] = {"room": {"Keuken": seeded["room"]["id"], "Badkamer": None}}
    result = client.post("/api/import/preview", json=body).json()
    assert result["valid_count"] == 2
    assert result["unmatched"] == {"category": ["Vloeren"], "room": [], "supplier": ["Tegelhuis"]}


def test_commit_creates_everything_in_one_go(client: TestClient, seeded: dict, data_dir: Path):
    upload = _dutch_upload(client, seeded)
    body = {
        "token": upload["token"],
        "mapping": upload["mapping"],
        "default_vat_rate": 21,
        "value_map": {"room": {"Keuken": seeded["room"]["id"]}},
    }
    result = client.post("/api/import/commit", json=body).json()
    assert result == {
        "created": 3,
        "skipped": 2,
        "payments": 2,
        "created_categories": 4,  # Ovens (under Kitchen), Vloeren, Vloeren > Tegels, Sanitair
        "created_rooms": 1,  # Badkamer, once for both spellings
        "created_suppliers": 1,  # Tegelhuis
    }
    assert list((data_dir / "imports").iterdir()) == []

    categories = {c["name"]: c for c in client.get("/api/categories").json()}
    assert categories["Ovens"]["parent_id"] == seeded["kitchen"]["id"]
    assert categories["Tegels"]["parent_id"] == categories["Vloeren"]["id"]
    assert categories["Sanitair"]["parent_id"] is None
    rooms = [r["name"] for r in client.get("/api/rooms").json()]
    assert sorted(rooms) == ["Badkamer", "Ground floor"]

    expenses = {e["title"]: e for e in client.get("/api/expenses").json()}
    assert sorted(expenses) == ["Kraan", "Oven", "Tegels"]
    oven = expenses["Oven"]
    assert oven["category_id"] == categories["Ovens"]["id"]
    assert oven["room_id"] == seeded["room"]["id"]
    assert oven["supplier_id"] == seeded["supplier"]["id"]
    assert oven["quote_ref"] == "OF-1"
    assert oven["created_by"] == "Tester"
    tiles = expenses["Tegels"]
    assert tiles["status"] == "paid"
    assert tiles["invoice_no"] == "F-22"
    assert tiles["notes"] == "met plinten"
    detail = client.get(f"/api/expenses/{tiles['id']}").json()
    assert detail["payments"][0]["reference"] == "Imported"
    assert detail["payments"][0]["method"] == "bank_transfer"
    assert detail["payments"][0]["date"] == "2026-02-15"
    tap = expenses["Kraan"]
    assert tap["paid_cents"] == 5_000
    assert tap["due_date"] == "2026-04-01"
    assert tap["room_id"] == expenses["Tegels"]["room_id"]

    # The upload is gone, so a second commit cannot import twice.
    assert client.post("/api/import/commit", json=body).status_code == 404


def test_commit_rolls_back_on_failure(client: TestClient, seeded: dict, monkeypatch):
    from app.services import importer

    upload = _dutch_upload(client, seeded)
    body = {"token": upload["token"], "mapping": upload["mapping"]}
    calls = []

    def failing(db, payload, user):
        calls.append(payload.title)
        if len(calls) == 2:
            raise RuntimeError("boom")
        return original(db, payload, user)

    original = importer.create_expense
    monkeypatch.setattr(importer, "create_expense", failing)
    with pytest.raises(RuntimeError):
        client.post("/api/import/commit", json=body)
    assert client.get("/api/expenses").json() == []
    assert len(client.get("/api/categories").json()) == 3
    assert client.post("/api/import/preview", json=body).status_code == 200


def test_csv_with_semicolons_and_comma_decimals(client: TestClient):
    text = (
        "Omschrijving;Bedrag incl. btw;BTW;Leverancier;Datum\n"
        "Verf;€ 1.210,00;21;Brico;03/01/2026\n"
        '"Lijm; extra sterk";12,10;21%;Brico;4.1.2026\n'
        ";;;;\n"
    )
    upload = _upload(client, text.encode("cp1252"), "export.csv")
    assert upload["sheets"] == []
    assert upload["sheet"] is None
    assert upload["row_count"] == 2
    assert upload["sample"][1][0] == "Lijm; extra sterk"
    assert upload["mapping"]["amount"] == 1
    assert upload["mapping"]["vat_rate"] == 2

    body = {"token": upload["token"], "mapping": upload["mapping"], "amount_includes_vat": True}
    preview = client.post("/api/import/preview", json=body).json()
    paint = preview["rows"][0]["values"]
    assert paint["amount_incl_cents"] == 121_000
    assert paint["amount_excl_cents"] == 100_000
    assert paint["stage"] == "invoiced"
    assert preview["rows"][1]["values"]["expense_date"] == "2026-01-04"

    result = client.post("/api/import/commit", json=body).json()
    assert result["created"] == 2
    assert result["created_suppliers"] == 1
    expenses = client.get("/api/expenses").json()
    assert sorted(e["amount_incl_cents"] for e in expenses) == [1_210, 121_000]


def test_value_map_links_name_to_existing_record(client: TestClient, seeded: dict):
    content = b"Title,Amount,Supplier,Category\nBoiler,1000.50,ACME bvba,Heating\n"
    upload = _upload(client, content, "boiler.csv")
    body = {
        "token": upload["token"],
        "mapping": upload["mapping"],
        "value_map": {
            "supplier": {"ACME bvba": seeded["supplier"]["id"]},
            "category": {"heating": seeded["electrical"]["id"]},
        },
    }
    preview = client.post("/api/import/preview", json=body).json()
    assert preview["unmatched"] == {"category": [], "room": [], "supplier": []}
    result = client.post("/api/import/commit", json=body).json()
    assert result["created_suppliers"] == 0
    assert result["created_categories"] == 0
    expense = client.get("/api/expenses").json()[0]
    assert expense["supplier_id"] == seeded["supplier"]["id"]
    assert expense["category_id"] == seeded["electrical"]["id"]
    assert expense["amount_excl_cents"] == 100_050

    bad = {**body, "value_map": {"room": {"X": 999}}}
    upload = _upload(client, content, "boiler.csv")
    bad["token"] = upload["token"]
    assert client.post("/api/import/preview", json=bad).status_code == 422


def test_template_round_trip(client: TestClient):
    response = client.get("/api/import/template")
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.content))
    assert wb.sheetnames == ["Expenses", "Instructions"]
    fields = client.get("/api/import/fields").json()
    assert [f["key"] for f in fields][:2] == ["title", "amount"]
    assert {f["key"]: f["kind"] for f in fields}["stage"] == "stage"
    headers = [c.value for c in wb["Expenses"][1]]
    assert headers == [f["label"] for f in fields]

    upload = _upload(client, response.content, "template.xlsx")
    assert upload["mapping"] == {f["key"]: i for i, f in enumerate(fields)}
    preview = client.post(
        "/api/import/preview", json={"token": upload["token"], "mapping": upload["mapping"]}
    ).json()
    assert preview["error_count"] == 0
    assert preview["rows"][0]["values"]["vat_rate"] == 21
